"""Unit tests for attachment and document file parsing."""

import io
import unittest
from locallm.core.file_parser import (
    extract_text_from_docx,
    extract_text_from_pdf,
    parse_attachment_file,
)


class TestFileParser(unittest.TestCase):
    """Test extracting text from various document formats."""

    def test_parse_plain_text(self):
        content = "Line 1\nLine 2\nLine 3"
        ok, text, err = parse_attachment_file("notes.txt", content.encode("utf-8"))
        self.assertTrue(ok)
        self.assertEqual(text, content)
        self.assertEqual(err, "")

    def test_parse_docx_with_document(self):
        import docx
        doc = docx.Document()
        doc.add_heading("Judul Dokumen", level=1)
        doc.add_paragraph("Paragraf pertama laporan.")
        doc.add_paragraph("Paragraf kedua laporan.")
        buf = io.BytesIO()
        doc.save(buf)
        raw_docx = buf.getvalue()

        ok, text, err = parse_attachment_file("dokumen.docx", raw_docx)
        self.assertTrue(ok)
        self.assertIn("Judul Dokumen", text)
        self.assertIn("Paragraf pertama laporan.", text)
        self.assertEqual(err, "")

    def test_parse_docx_fallback(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("Isi teks fallback.")
        buf = io.BytesIO()
        doc.save(buf)
        raw_docx = buf.getvalue()

        # Direct test on extract_text_from_docx
        extracted = extract_text_from_docx(raw_docx)
        self.assertIn("Isi teks fallback.", extracted)

    def test_unsupported_or_empty(self):
        ok, text, err = parse_attachment_file("kosong.docx", b"not-a-valid-zip-or-docx")
        self.assertFalse(ok)
        self.assertIn("Failed to parse", err)


if __name__ == "__main__":
    unittest.main()
