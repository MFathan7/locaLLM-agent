"""Document and attachment parser for extracting clean text from various formats (.docx, .pdf, txt, etc.)."""

import io
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Optional, Union, Tuple


def extract_text_from_docx(file_source: Union[str, Path, bytes, io.BytesIO]) -> str:
    """Extract clean text content from a .docx file using python-docx or standard zipfile fallback."""
    # Attempt 1: Try using python-docx if installed
    try:
        import docx  # type: ignore
        if isinstance(file_source, (str, Path)):
            doc = docx.Document(file_source)
        elif isinstance(file_source, bytes):
            doc = docx.Document(io.BytesIO(file_source))
        else:
            doc = docx.Document(file_source)

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_texts.append(" | ".join(row_cells))
        
        combined = []
        if paragraphs:
            combined.extend(paragraphs)
        if table_texts:
            combined.append("\n[Tables Content:]")
            combined.extend(table_texts)
        
        return "\n\n".join(combined).strip()
    except Exception:
        pass

    # Attempt 2: Fallback to built-in zipfile + XML extraction (zero dependencies)
    try:
        if isinstance(file_source, (str, Path)):
            z = zipfile.ZipFile(file_source)
        elif isinstance(file_source, bytes):
            z = zipfile.ZipFile(io.BytesIO(file_source))
        else:
            z = zipfile.ZipFile(file_source)

        with z:
            xml_content = z.read("word/document.xml")
        
        tree = ET.fromstring(xml_content)
        # WordML namespace
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        
        paragraphs = []
        for p in tree.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
            p_text = "".join(
                t.text for t in p.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t") if t.text
            )
            if p_text.strip():
                paragraphs.append(p_text.strip())
        
        return "\n\n".join(paragraphs).strip()
    except Exception as exc:
        raise ValueError(f"Failed to parse .docx file: {exc}")


def extract_text_from_pdf(file_source: Union[str, Path, bytes, io.BytesIO]) -> str:
    """Extract clean text content from a .pdf file using pypdf."""
    try:
        import pypdf  # type: ignore
        if isinstance(file_source, (str, Path)):
            reader = pypdf.PdfReader(file_source)
        elif isinstance(file_source, bytes):
            reader = pypdf.PdfReader(io.BytesIO(file_source))
        else:
            reader = pypdf.PdfReader(file_source)

        extracted_pages = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                extracted_pages.append(f"--- Page {idx + 1} ---\n{text.strip()}")

        return "\n\n".join(extracted_pages).strip()
    except Exception as exc:
        raise ValueError(f"Failed to parse .pdf file: {exc}")


def parse_attachment_file(
    filename: str,
    raw_bytes: bytes,
    max_chars: int = 150000,
) -> Tuple[bool, str, str]:
    """Parse raw bytes of an attachment into clean readable text based on filename extension.

    Returns:
        (success: bool, text_content: str, error_message: str)
    """
    ext = Path(filename).suffix.lower()

    try:
        # Word Documents
        if ext in (".docx", ".docm", ".dotx"):
            text = extract_text_from_docx(raw_bytes)
            if not text:
                return False, "", "File .docx tidak memiliki teks yang dapat diekstrak atau kosong."
            if len(text) > max_chars:
                text = text[:max_chars] + f"\n\n[... Truncated: document text exceeded {max_chars} characters ...]"
            return True, text, ""

        # PDF Documents
        if ext == ".pdf":
            text = extract_text_from_pdf(raw_bytes)
            if not text:
                return False, "", "File PDF tidak memiliki teks yang dapat diekstrak atau merupakan pindaian gambar (scanned)."
            if len(text) > max_chars:
                text = text[:max_chars] + f"\n\n[... Truncated: PDF text exceeded {max_chars} characters ...]"
            return True, text, ""

        # Plain text / code / CSV / JSON
        try:
            text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = raw_bytes.decode("latin-1")
            except Exception:
                return False, "", f"File '{filename}' tidak dapat dibaca sebagai format teks yang valid."

        if len(text) > max_chars:
            text = text[:max_chars] + f"\n\n[... Truncated: text exceeded {max_chars} characters ...]"
        return True, text, ""

    except Exception as exc:
        return False, "", str(exc)
